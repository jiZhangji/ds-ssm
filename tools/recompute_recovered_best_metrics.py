import csv
import json
import os
from glob import glob

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


MANIFEST = "downloads/experiment_recovery/recovery_manifest.json"
OUT_JSON = "database/experiment_ranking_recomputed.json"
OUT_CSV = "database/experiment_ranking_recomputed.csv"
OUT_XLSX = "database/experiment_ranking_recomputed.xlsx"
OUT_MD = "reports/experiment_ranking_recomputed.md"
SUMMARY_XLSX = "database/SSDD_experiment_results_summary_recomputed.xlsx"
SUMMARY_CSV = "database/SSDD_experiment_results_summary_recomputed.csv"


def fnum(v):
    try:
        return float(str(v).strip())
    except Exception:
        return None


def find_col(fieldnames, *needles):
    for n in needles:
        for f in fieldnames:
            if f and n.lower() == f.strip().lower():
                return f
    for n in needles:
        for f in fieldnames:
            if f and n.lower() in f.strip().lower():
                return f
    return None


def best_from_results(path):
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        fields = reader.fieldnames or []
        epoch_col = find_col(fields, "epoch")
        p_col = find_col(fields, "metrics/precision(B)", "precision")
        r_col = find_col(fields, "metrics/recall(B)", "recall")
        map50_col = find_col(fields, "metrics/mAP50(B)", "mAP50(B)")
        map_col = find_col(fields, "metrics/mAP50-95(B)", "mAP50-95(B)")
        if not map_col:
            return None
        best = None
        rows = 0
        for row in reader:
            rows += 1
            score = fnum(row.get(map_col))
            if score is None:
                continue
            if best is None or score > best["mAP50-95"]:
                best = {
                    "best_epoch": fnum(row.get(epoch_col)),
                    "P": fnum(row.get(p_col)),
                    "R": fnum(row.get(r_col)),
                    "mAP50": fnum(row.get(map50_col)),
                    "mAP50-95": score,
                    "epochs_recorded": rows,
                    "results_csv": path.replace("\\", "/"),
                }
        return best


def main():
    os.makedirs("database", exist_ok=True)
    os.makedirs("reports", exist_ok=True)
    manifest = json.load(open(MANIFEST, encoding="utf-8"))
    rows = []
    for item in manifest:
        result_paths = [d.get("local_path") for d in item.get("downloaded", []) if d.get("local_path", "").endswith("results.csv")]
        if not result_paths:
            continue
        best = best_from_results(result_paths[0])
        if not best:
            continue
        row = {
            "experiment_name": item["experiment"],
            "source_server": item["source_server"],
            "host": item["host"],
            "remote_dir": item["remote_dir"],
            "old_rank": item.get("rank"),
            "old_score": item.get("score"),
            **best,
        }
        rows.append(row)
    rows.sort(key=lambda r: r["mAP50-95"], reverse=True)
    for i, row in enumerate(rows, 1):
        row["rank"] = i

    fields = ["rank", "experiment_name", "source_server", "best_epoch", "P", "R", "mAP50", "mAP50-95", "epochs_recorded", "old_rank", "old_score", "remote_dir", "results_csv"]
    with open(OUT_CSV, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows([{k: r.get(k, "") for k in fields} for r in rows])
    with open(OUT_JSON, "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=2)

    wb = Workbook()
    ws = wb.active
    ws.title = "recomputed_ranking"
    ws.append(fields)
    for r in rows:
        ws.append([r.get(k, "") for k in fields])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for col in range(1, ws.max_column + 1):
        max_len = max(len(str(ws.cell(r, col).value or "")) for r in range(1, min(ws.max_row, 80) + 1))
        ws.column_dimensions[get_column_letter(col)].width = min(max(max_len + 2, 10), 60)
    wb.save(OUT_XLSX)

    md = ["# Recomputed Experiment Ranking", "", "Ranking source: recovered local `results.csv` files. Score is max `metrics/mAP50-95(B)` across epochs.", "", "| Rank | Server | Experiment | Best epoch | P | R | mAP50 | mAP50-95 | Old score |", "|---:|---|---|---:|---:|---:|---:|---:|---:|"]
    for r in rows:
        md.append(f"| {r['rank']} | `{r['source_server']}` | `{r['experiment_name']}` | `{r['best_epoch']}` | `{r['P']}` | `{r['R']}` | `{r['mAP50']}` | `{r['mAP50-95']}` | `{r.get('old_score','')}` |")
    open(OUT_MD, "w", encoding="utf-8").write("\n".join(md) + "\n")

    # Recomputed SSDD summary focused on actual recovered results.
    sum_fields = ["实验序号", "实验名称", "来源服务器", "最佳epoch", "P", "R", "mAP50", "mAP50-95", "原rank", "原score", "实际实验路径", "本地results.csv"]
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.title = "SSDD重算结果总表"
    ws2.append(sum_fields)
    with open(SUMMARY_CSV, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(sum_fields)
        for r in rows:
            vals = [r["rank"], r["experiment_name"], r["source_server"], r["best_epoch"], r["P"], r["R"], r["mAP50"], r["mAP50-95"], r["old_rank"], r["old_score"], r["remote_dir"], r["results_csv"]]
            ws2.append(vals)
            w.writerow(vals)
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    ws2.freeze_panes = "A2"
    for col in range(1, ws2.max_column + 1):
        max_len = max(len(str(ws2.cell(r, col).value or "")) for r in range(1, min(ws2.max_row, 80) + 1))
        ws2.column_dimensions[get_column_letter(col)].width = min(max(max_len + 2, 10), 60)
    wb2.save(SUMMARY_XLSX)
    print(f"recomputed {len(rows)} experiments")


if __name__ == "__main__":
    main()
