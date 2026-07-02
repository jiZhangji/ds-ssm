from pathlib import Path
import sys

sys.path.insert(0, r"C:\Users\13207\.cache\codex-runtimes\codex-primary-runtime\dependencies\python")

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def set_east_asian_font(run, font_name="宋体"):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def set_doc_defaults(doc):
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(10.5)


def add_heading(doc, text, level=1):
    p = doc.add_heading("", level=level)
    run = p.add_run(text)
    set_east_asian_font(run, "黑体")
    return p


def shade_header(row):
    for cell in row.cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "D9EAF7")
        tc_pr.append(shd)


def main():
    root = Path.cwd()
    out = root / "paper_draft" / "DS_SSM_方法优势与MaDiNet对比说明_中文修正版.docx"
    out.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    set_doc_defaults(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("DS-SSM 方法优势与 MaDiNet 对比叙事建议")
    run.bold = True
    run.font.size = Pt(16)
    set_east_asian_font(run, "黑体")

    doc.add_paragraph("生成日期：2026-06-14")
    doc.add_paragraph("说明：本文件是论文写作草稿，不代表已经完成 MaDiNet 官方代码复现实验。MaDiNet 的参数量、FLOPs 和速度需要后续在服务器上统一测量或从论文中严格引用。")

    add_heading(doc, "一、需要先澄清的点", 1)
    doc.add_paragraph(
        "当前我没有在服务器上实际运行 MaDiNet 或本方法的复杂度实验。上一版生成得很快，是因为只完成了本地文档和脚本整理。"
        "因此，文中关于 MaDiNet 的参数量、FLOPs、FPS 不能写成已测结果，只能作为待测项。"
    )
    doc.add_paragraph(
        "此外，上一版 Word 出现乱码，是因为通过 PowerShell 管道向 Python 传递中文时发生编码损坏。本修正版使用 UTF-8 脚本直接生成，中文应能正常显示。"
    )

    add_heading(doc, "二、当前论文背景建议", 1)
    doc.add_paragraph(
        "SSDD 数据集中的 SAR 舰船目标通常具有尺度小、背景杂波强、斑点噪声明显、目标边界不稳定等特点。"
        "MaDiNet 等方法在 SSDD 上已经取得了很强的检测性能，说明单纯追求更高 mAP50 的空间会变小。"
        "因此，本工作的论文叙事应从“高效 SAR 舰船检测”角度展开：在保持接近强参考方法检测性能的同时，显著降低参数量和计算量。"
    )
    doc.add_paragraph(
        "这种定位对论文是有利的。因为实际 SAR 图像处理往往面向大范围海域、批量遥感数据和部署资源受限场景，"
        "模型不仅要检得准，还要足够轻、足够快、足够稳定。"
    )

    add_heading(doc, "三、建议融入论文的创新点", 1)
    items = [
        (
            "DS-SSM 是核心创新",
            "DS-SSM（De-Speckle Selective State Space Model）应作为本文方法的核心。它面向 SAR 斑点噪声和海杂波背景，"
            "在选择性状态空间建模中引入去斑感知和目标响应保持机制，避免状态传播过程中过度积累背景噪声。"
        ),
        (
            "PFIM 是 DS-SSM 内部的重要增强",
            "PFIM 可以表述为感知重要性引导模块。它不是完全脱离 DS-SSM 的另一套方法，而是服务于 DS-SSM 的状态更新，"
            "通过重要性图调制 gamma、write gate 或 delta，使模型更关注疑似舰船区域。"
        ),
        (
            "DG-DS-SSM 是延长线",
            "DG-DS-SSM 可以作为进一步增强方向。它把状态更新拆成 erase、write 和 delta 等更细粒度门控，"
            "适合写成方法拓展或消融实验，而不是抢走 DS-SSM 的主创新位置。"
        ),
        (
            "复杂度优势是论文的重要卖点",
            "如果 MaDiNet 在 SSDD 上的最高精度仍然更强，本方法仍然可以通过更低参数量、更低 FLOPs 和更快推理速度形成优势。"
            "论文中应加入 Params、GFLOPs、FPS、Latency 等指标，构建精度-效率联合对比。"
        ),
    ]
    for title, body in items:
        p = doc.add_paragraph()
        r = p.add_run(title + "：")
        r.bold = True
        set_east_asian_font(r, "黑体")
        p.add_run(body)

    add_heading(doc, "四、与 MaDiNet 的合理对比方式", 1)
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["对比维度", "MaDiNet", "本文方法", "论文表达建议"]
    for i, text in enumerate(headers):
        table.rows[0].cells[i].text = text
    shade_header(table.rows[0])

    rows = [
        ["检测精度", "SSDD 上表现很强，可能难以直接超越", "mAP50 接近强参考方法，P/R 表现稳定", "承认其强性能，同时强调本文接近强基线"],
        ["参数量", "待官方代码或论文数值确认", "预期更小，需要统一脚本实测", "突出轻量化"],
        ["FLOPs", "待测；若包含复杂推理过程，实际代价可能更高", "单阶段检测框架，计算链路更直接", "突出效率-精度平衡"],
        ["推理速度", "需在相同 GPU 上统一测试", "可在 3090 上统计 FPS 和 latency", "强调实际部署价值"],
        ["论文定位", "强性能参考方法", "高效去斑选择性状态空间检测方法", "避免只做单一 mAP 竞赛"],
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text

    add_heading(doc, "五、当前实验结果应如何写", 1)
    doc.add_paragraph(
        "当前已经整理的本项目结果中，PFIM / IG-DS-SSM 主线代表性结果为："
        "P=0.97378，R=0.95219，mAP50=0.98704，mAP50-95=0.74022。"
        "DG-DS-SSM 延长线在 TTA IoU=0.55 设置下达到：P=0.963，R=0.962，mAP50=0.987，mAP50-95=0.724。"
    )
    doc.add_paragraph(
        "考虑 SSDD 数据集的特点，论文中建议把 mAP50、Precision、Recall 作为主要指标，mAP50-95 作为辅助指标。"
        "同时必须补充参数量和 FLOPs 表格，否则“轻量高效”的优势缺少硬证据。"
    )

    add_heading(doc, "六、可直接放入论文的中文贡献表述", 1)
    contributions = [
        "提出 DS-SSM，一种面向 SAR 舰船检测的去斑选择性状态空间模块，用于缓解斑点噪声和海杂波在状态传播过程中的干扰。",
        "设计感知重要性引导机制 PFIM，将目标相关重要性图注入状态更新过程，使模型在保持长程建模能力的同时更加关注舰船候选区域。",
        "在 Mamba-YOLO 框架中构建轻量化 SAR 舰船检测模型，在 SSDD 上获得具有竞争力的 mAP50、Precision 和 Recall，同时具备更低参数量和 FLOPs 的潜力。",
        "进一步探索 DG-DS-SSM 延展结构，将状态更新解耦为擦除、写入和步长调制，为后续 SAR 状态空间模型设计提供可扩展方向。",
    ]
    for item in contributions:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "七、下一步必须补的实验", 1)
    experiments = [
        "在同一服务器、同一输入尺寸 640、同一 batch 设置下统计本文模型的 Params、GFLOPs、Latency 和 FPS。",
        "获取 MaDiNet 官方代码后，用同一输入尺寸统计其复杂度；如果无法运行官方代码，则严格引用论文报告值并标注 paper reported。",
        "补充精度-复杂度联合表格：Method、Params(M)、GFLOPs、FPS、P、R、mAP50、mAP50-95。",
        "如 MaDiNet 是扩散式或多步推理方法，需要额外记录实际推理延迟，因为静态 FLOPs 可能不能完全反映真实计算成本。",
    ]
    for item in experiments:
        doc.add_paragraph(item, style="List Number")

    add_heading(doc, "八、已经准备好的复杂度实验脚本", 1)
    paths = [
        root / "complexity_benchmark" / "model_complexity_benchmark.py",
        root / "complexity_benchmark" / "run_complexity_benchmark.sh",
        root / "complexity_benchmark" / "madinet_record_template.csv",
        root / "reports" / "madinet_complexity_experiment_plan.md",
    ]
    for path in paths:
        doc.add_paragraph(str(path), style="List Bullet")

    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
