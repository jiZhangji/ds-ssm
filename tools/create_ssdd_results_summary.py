import os

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


INPUT = r"D:\A-主机文件\Codex\Mamba-YOLO-ds-ssm\database\SSDD_experiment_data_structured_settings_verified.xlsx"
OUT_XLSX = r"D:\A-主机文件\Codex\Mamba-YOLO-ds-ssm\database\SSDD_experiment_results_summary.xlsx"
OUT_CSV = r"D:\A-主机文件\Codex\Mamba-YOLO-ds-ssm\database\SSDD_experiment_results_summary.csv"

SOURCE_SHEET = "结构化总表_核验更新"

KEEP_COLUMNS = [
    "序号",
    "实验编号",
    "实验名_raw",
    "实验阶段/别名",
    "变体/推理版本",
    "方法大类",
    "模型架构",
    "模型尺寸",
    "epochs",
    "batch_size",
    "DDP/多卡",
    "DS-SSM是否启用",
    "DS-SSM版本",
    "放置_raw",
    "lambda_init",
    "gamma是否启用",
    "gamma模式",
    "gamma范围",
    "PFIM是否启用",
    "PFIM_std类型",
    "PFIM_on_gamma",
    "PFIM_on_write_gate",
    "PFIM_on_delta",
    "滤波/空间模块",
    "滤波参数冻结",
    "推理模式",
    "TTA",
    "imgsz",
    "iou",
    "conf",
    "P",
    "R",
    "mAP50",
    "mAP50-95",
    "F1",
    "实际来源服务器",
    "实际rank",
    "实际score",
    "实际实验路径",
    "匹配状态",
    "核验备注",
    "是否推荐论文主表",
    "论文主表角色/推荐用途",
    "异常/注意事项",
    "配置继承/解析说明",
    "原始方法与主要改动",
]


def val(ws, row, headers, name):
    col = headers.get(name)
    return ws.cell(row, col).value if col else None


def main():
    os.makedirs(os.path.dirname(OUT_XLSX), exist_ok=True)
    src_wb = openpyxl.load_workbook(INPUT, data_only=True)
    ws = src_wb[SOURCE_SHEET]
    headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}

    wb = Workbook()
    out = wb.active
    out.title = "SSDD实验结果总表"
    out.append(KEEP_COLUMNS)

    rows = []
    for r in range(2, ws.max_row + 1):
        row = [val(ws, r, headers, c) for c in KEEP_COLUMNS]
        rows.append(row)

    # Sort by actual rank when present, otherwise original sequence.
    rank_idx = KEEP_COLUMNS.index("实际rank")
    seq_idx = KEEP_COLUMNS.index("序号")
    rows.sort(key=lambda x: (x[rank_idx] is None, x[rank_idx] if x[rank_idx] is not None else 999999, x[seq_idx] or 999999))

    for row in rows:
        out.append(row)

    for cell in out[1]:
        cell.font = Font(bold=True)
    out.freeze_panes = "A2"
    for col in range(1, out.max_column + 1):
        max_len = 0
        for row in range(1, min(out.max_row, 80) + 1):
            max_len = max(max_len, len(str(out.cell(row, col).value or "")))
        out.column_dimensions[get_column_letter(col)].width = min(max(max_len + 2, 10), 48)

    wb.save(OUT_XLSX)

    import csv

    with open(OUT_CSV, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(KEEP_COLUMNS)
        writer.writerows(rows)

    print(OUT_XLSX)
    print(OUT_CSV)


if __name__ == "__main__":
    main()
