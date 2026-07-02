import openpyxl

path = r"C:\Users\13207\Desktop\GPT-文档\SSDD_experiment_data_structured_settings.xlsx"
wb = openpyxl.load_workbook(path, data_only=False)
print(wb.sheetnames)
for ws in wb.worksheets:
    print("SHEET", ws.title, ws.max_row, ws.max_column)
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    print("HEADERS", headers)
    for r in range(2, min(ws.max_row, 5) + 1):
        print([ws.cell(r, c).value for c in range(1, min(ws.max_column, 10) + 1)])
