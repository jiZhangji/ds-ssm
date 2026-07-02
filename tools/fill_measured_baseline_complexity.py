from pathlib import Path
import csv
import json
import sys

sys.path.insert(0, r"C:\Users\13207\.cache\codex-runtimes\codex-primary-runtime\dependencies\python")

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path.cwd()
PLAN = ROOT / "database" / "open_source_ssdd_complexity_comparison_plan.csv"
YOLO_RESULTS = ROOT / "downloads" / "complexity_results_172_26_151_143" / "ultralytics_baselines_20260614_181605" / "ultralytics_baseline_complexity.csv"


MAP_NAMES = {
    "yolov3.yaml": "YOLOv3",
    "yolov5n.yaml": "YOLOv5n",
    "yolov5s.yaml": "YOLOv5s",
    "yolov8n.yaml": "YOLOv8n",
}


def read_csv(path):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_csv(path, rows, fields):
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def f1(p, r):
    p = float(p)
    r = float(r)
    return round(2 * p * r / (p + r), 3)


def set_font(run, font="宋体"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)


def heading(doc, text, level=1):
    p = doc.add_heading("", level=level)
    rr = p.add_run(text)
    set_font(rr, "黑体")


def shade(row):
    for cell in row.cells:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "D9EAF7")
        cell._tc.get_or_add_tcPr().append(shd)


def main():
    rows = read_csv(PLAN)
    yolo_rows = read_csv(YOLO_RESULTS)
    measured = {}
    for r in yolo_rows:
        method = MAP_NAMES.get(r["model"])
        if method:
            measured[method] = r
    for r in rows:
        if r["method"] in measured:
            m = measured[r["method"]]
            r["params_M"] = m["params_M"]
            r["GFLOPs"] = m["GFLOPs"]
            r["latency_ms"] = m["latency_ms_per_image"]
            r["fps"] = m["fps"]
            r["status"] = "measured on server 172.26.151.143 using Ultralytics config, input 640, batch=1, fp32"
        if r["method"] == "Ours: DS-SSM/PFIM":
            r["F1_percent_from_MaDiNet"] = str(f1(r["P_percent_from_MaDiNet"], r["R_percent_from_MaDiNet"]))
        if r["method"] == "Ours: DG-DS-SSM":
            r["F1_percent_from_MaDiNet"] = str(f1(r["P_percent_from_MaDiNet"], r["R_percent_from_MaDiNet"]))
    fields = list(rows[0].keys())
    write_csv(PLAN, rows, fields)

    md = ROOT / "reports" / "open_source_ssdd_complexity_comparison_plan.md"
    lines = [
        "# SSDD 开源方法精度-复杂度对比计划",
        "",
        "生成时间：2026-06-14",
        "",
        "本表只包含 MaDiNet SSDD 对比表中标注为开源的方法，以及本文方法。非开源方法不进入主复现实验表，只作为论文背景引用。",
        "",
        "| 方法 | AP50 | P | R | F1 | Params(M) | GFLOPs | Latency(ms) | FPS | 状态 |",
        "|---|---:|---:|---:|---:|---:|---:|---:|---:|---|",
    ]
    for r in rows:
        lines.append(
            f"| {r['method']} | {r['AP50_percent_from_MaDiNet']} | {r['P_percent_from_MaDiNet']} | {r['R_percent_from_MaDiNet']} | {r['F1_percent_from_MaDiNet']} | {r['params_M']} | {r['GFLOPs']} | {r['latency_ms']} | {r['fps']} | {r['status']} |"
        )
    lines += [
        "",
        "## 已实测",
        "",
        "- YOLOv3、YOLOv5n、YOLOv5s、YOLOv8n 已在服务器 `172.26.151.143` 上完成复杂度统计。",
        "- DS-SSM/PFIM、DG-DS-SSM 已完成复杂度统计。",
        "- MaDiNet 官方 MambaSAR res4+res5 核心模块已完成复杂度统计，但完整 detector 仍待 Detectron2 环境。",
        "",
        "## 当前无法直接测量的原因",
        "",
        "服务器当前环境缺少 `mmdet`、`mmengine`、`mmcv`、`yolox`、`detectron2`，因此 MMDetection/YOLOX/DiffusionDet/ConsistencyDet/DiffDet4SAR/MaDiNet 完整 detector 暂不能直接构建统计。",
    ]
    md.write_text("\n".join(lines), encoding="utf-8")

    doc_path = ROOT / "paper_draft" / "DS_SSM_顶会论文背景动机与轻量化对比框架.docx"
    doc = Document(str(doc_path))
    heading(doc, "八、已补测的开源 baseline 复杂度", 1)
    doc.add_paragraph("已在服务器 172.26.151.143 上补测当前环境可直接构建的 Ultralytics 开源 baseline。")
    table = doc.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["方法", "Params(M)", "GFLOPs", "Latency(ms)", "FPS", "说明"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    shade(table.rows[0])
    for method in ["YOLOv3", "YOLOv5n", "YOLOv5s", "YOLOv8n", "Ours: DS-SSM/PFIM", "Ours: DG-DS-SSM", "MaDiNet"]:
        r = next(x for x in rows if x["method"] == method)
        cells = table.add_row().cells
        vals = [r["method"], r["params_M"], r["GFLOPs"], r["latency_ms"], r["fps"], r["status"]]
        for i, v in enumerate(vals):
            cells[i].text = str(v)
    doc.add_paragraph("剩余 TO_MEASURE 项不是不测，而是当前服务器环境缺少对应框架；需要新建 MMDetection/YOLOX/Detectron2 环境后继续补齐。")
    doc.save(doc_path)

    print(json.dumps({"updated": str(PLAN), "report": str(md), "docx": str(doc_path)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
