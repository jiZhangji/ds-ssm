from pathlib import Path
import csv
import json
import sys

sys.path.insert(0, r"C:\Users\13207\.cache\codex-runtimes\codex-primary-runtime\dependencies\python")

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path.cwd()
OURS = ROOT / "downloads" / "complexity_results_172_26_151_143" / "server_run_20260614_171404" / "ours_complexity.csv"
MADINET = ROOT / "downloads" / "complexity_results_172_26_151_143" / "madinet_component_20260614_171712" / "madinet_official_complexity.csv"


def read_csv(path):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_csv(path, rows, fields):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def set_font(run, name="宋体"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def heading(doc, text, level=1):
    p = doc.add_heading("", level=level)
    r = p.add_run(text)
    set_font(r, "黑体")
    return p


def shade(row):
    for cell in row.cells:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "D9EAF7")
        cell._tc.get_or_add_tcPr().append(shd)


def make_outputs():
    ours = read_csv(OURS)
    madinet = read_csv(MADINET)[0]
    rows = []
    for r in ours:
        rows.append(
            {
                "method": r["model"].replace(".yaml", ""),
                "scope": "full Mamba-YOLO detector",
                "params_M": r["params_M"],
                "GFLOPs": r["GFLOPs"],
                "latency_ms": r["latency_ms_per_image"],
                "fps": r["fps"],
                "device": r["device_name"],
                "source": "server measured",
                "note": "640 input, batch=1, fp32",
            }
        )
    rows.append(
        {
            "method": "MaDiNet official MambaSAR res4+res5",
            "scope": "official core module only; not full detector",
            "params_M": madinet["params_M"],
            "GFLOPs": madinet["GFLOPs_thop"],
            "latency_ms": madinet["latency_ms"],
            "fps": "",
            "device": madinet["device_name"],
            "source": "official code measured",
            "note": "Official MambaSAR.py res4/res5 block; full detector requires Detectron2 config/environment.",
        }
    )
    fields = ["method", "scope", "params_M", "GFLOPs", "latency_ms", "fps", "device", "source", "note"]
    write_csv(ROOT / "database" / "complexity_comparison_madinet_ds_ssm.csv", rows, fields)

    md = ROOT / "reports" / "madinet_ds_ssm_complexity_report.md"
    md.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "# MaDiNet 与 DS-SSM 复杂度实测报告",
        "",
        "生成日期：2026-06-14",
        "",
        "## 服务器与环境",
        "",
        "- 服务器：172.26.151.143",
        "- 项目目录：`/mnt/data5/zhangji/Mamba-YOLO`",
        "- GPU：NVIDIA GeForce RTX 3090",
        "- PyTorch：2.3.0+cu121",
        "- 输入尺寸：640",
        "- batch：1",
        "- 精度：fp32",
        "",
        "## 实测结果",
        "",
        "| 方法 | 范围 | Params(M) | GFLOPs | latency(ms/img) | FPS | 说明 |",
        "|---|---|---:|---:|---:|---:|---|",
    ]
    for r in rows:
        lines.append(
            f"| {r['method']} | {r['scope']} | {r['params_M']} | {r['GFLOPs']} | {r['latency_ms']} | {r['fps']} | {r['note']} |"
        )
    lines += [
        "",
        "## 重要说明",
        "",
        "1. 本项目三种 Mamba-YOLO 检测器均已在服务器上实际构建并统计，结果可作为论文复杂度表的直接依据。",
        "2. MaDiNet 官方仓库已经下载并上传到服务器，统计的是官方 `MambaSAR.py` 中 res4+res5 核心模块，不是完整 Detectron2 检测器。",
        "3. 完整 MaDiNet 检测器暂未 end-to-end 统计，原因是服务器当前环境缺少 `detectron2`，且官方仓库未提供清晰的 SSDD 专用配置文件入口。",
        "4. 仅 MaDiNet 核心 MambaSAR res4+res5 模块已经达到 223.26M 参数和 345.28 GFLOPs，显著大于本文完整检测器约 5.96M 参数和 13.16-13.21 GFLOPs。",
        "5. 因 MaDiNet 使用自定义 selective scan/CUDA 路径，THOP 可能低估或不完全覆盖部分算子，论文中应将该数值标注为基于官方模块的 profiler 结果。",
        "",
        "## 论文结论建议",
        "",
        "本文方法不必只围绕是否超过 MaDiNet 的 SSDD 最高 mAP 展开，而应强调在接近强参考方法精度的同时，模型规模和计算量显著降低。DS-SSM / PFIM 的优势是以轻量结构实现 SAR 去斑选择性建模，适合实际部署和后续扩展。",
    ]
    md.write_text("\n".join(lines), encoding="utf-8")

    doc = Document()
    doc.styles["Normal"].font.name = "宋体"
    doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    doc.styles["Normal"].font.size = Pt(10.5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DS-SSM 方法优势与 MaDiNet 复杂度实测对比")
    r.bold = True
    r.font.size = Pt(16)
    set_font(r, "黑体")

    doc.add_paragraph("生成日期：2026-06-14")
    doc.add_paragraph("说明：本文档基于服务器 172.26.151.143 上的真实复杂度统计结果生成。")

    heading(doc, "一、背景描述建议", 1)
    doc.add_paragraph(
        "MaDiNet 在 SSDD 数据集上具有很强检测性能，因此本文不应只把叙事建立在单一 mAP 超越上。"
        "更合理的论文定位是：面向 SAR 舰船检测的高效去斑选择性状态空间模型。"
        "在实际遥感应用中，模型参数量、FLOPs、推理速度和部署成本同样重要。"
    )
    doc.add_paragraph(
        "服务器实测显示，本文完整 DS-SSM / DG-DS-SSM 检测器约 5.96M 参数、13.16-13.21 GFLOPs；"
        "而 MaDiNet 官方核心 MambaSAR res4+res5 模块单独就达到 223.26M 参数和 345.28 GFLOPs。"
        "这说明本文方法在模型复杂度上具有非常明确的优势。"
    )

    heading(doc, "二、复杂度实测表", 1)
    table = doc.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["方法", "统计范围", "Params(M)", "GFLOPs", "延迟(ms/img)", "说明"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    shade(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        vals = [row["method"], row["scope"], row["params_M"], row["GFLOPs"], row["latency_ms"], row["note"]]
        for i, v in enumerate(vals):
            cells[i].text = str(v)

    heading(doc, "三、建议融入论文的创新点", 1)
    bullets = [
        "DS-SSM 是本文核心创新：它面向 SAR 斑点噪声和海杂波，将去斑感知融入选择性状态空间建模。",
        "PFIM 是 DS-SSM 的重要内部增强：通过目标重要性图调制状态更新，使模型更关注舰船候选区域。",
        "本文方法不是通过堆叠重型模块获得性能，而是在约 6M 参数、约 13G FLOPs 的轻量检测器中获得具有竞争力的 SSDD 表现。",
        "与 MaDiNet 这类强性能但结构复杂的方法相比，本文的突出优势是精度-效率平衡，更适合资源受限和大范围 SAR 图像处理场景。",
    ]
    for b in bullets:
        doc.add_paragraph(b, style="List Bullet")

    heading(doc, "四、可直接写入论文的贡献表述", 1)
    doc.add_paragraph(
        "We propose DS-SSM, an efficient de-speckle selective state space module for SAR ship detection. "
        "Unlike heavy Mamba-diffusion designs, DS-SSM introduces SAR-oriented noise-aware state modulation within a compact detector, "
        "achieving competitive SSDD accuracy with substantially fewer parameters and lower FLOPs."
    )
    doc.add_paragraph(
        "中文表述：本文提出 DS-SSM，一种面向 SAR 舰船检测的高效去斑选择性状态空间模块。"
        "不同于依赖重型 Mamba 或扩散结构的方法，DS-SSM 在紧凑检测器中引入噪声感知状态调制，"
        "在 SSDD 上获得具有竞争力的检测性能，同时显著降低参数量和计算量。"
    )

    heading(doc, "五、结果使用注意", 1)
    notes = [
        "MaDiNet 当前统计的是官方 MambaSAR 核心模块，不是完整检测器。",
        "完整 MaDiNet 需要 Detectron2 环境和明确 SSDD 配置后才能 end-to-end 统计。",
        "论文中可以先使用该结果说明 MaDiNet 核心模块的复杂度较高，但最终表格最好继续补充完整检测器的官方参数/FLOPs或论文报告值。",
    ]
    for n in notes:
        doc.add_paragraph(n, style="List Number")

    out_docx = ROOT / "paper_draft" / "DS_SSM_方法优势与MaDiNet复杂度实测对比_中文.docx"
    out_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_docx)

    summary = {
        "csv": str(ROOT / "database" / "complexity_comparison_madinet_ds_ssm.csv"),
        "report": str(md),
        "docx": str(out_docx),
        "rows": rows,
    }
    (ROOT / "reports" / "madinet_complexity_outputs.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    make_outputs()
