import json
import os
import re
from copy import copy
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


INPUT = r"C:\Users\13207\Desktop\GPT-文档\SSDD_experiment_data_structured_settings.xlsx"
OUTPUT = r"D:\A-主机文件\Codex\Mamba-YOLO-ds-ssm\database\SSDD_experiment_data_structured_settings_verified.xlsx"
RANKING = r"D:\A-主机文件\Codex\Mamba-YOLO-ds-ssm\database\experiment_ranking.json"


def norm(s):
    s = "" if s is None else str(s)
    return re.sub(r"[^a-z0-9]+", "", s.lower())


def score_fields(metrics):
    def pick(*names):
        for name in names:
            if name in metrics:
                return metrics[name]
        for k, v in metrics.items():
            kl = k.lower()
            if any(n.lower() in kl for n in names):
                return v
        return None

    return {
        "P": pick("max_metrics/precision(B)", "metrics/precision(B)", "precision"),
        "R": pick("max_metrics/recall(B)", "metrics/recall(B)", "recall"),
        "mAP50": pick("max_metrics/mAP50(B)", "metrics/mAP50(B)", "map50"),
        "mAP50-95": pick("max_metrics/mAP50-95(B)", "metrics/mAP50-95(B)", "map50-95"),
    }


def manual_to_query(raw, exp_id):
    text = f"{raw or ''} {exp_id or ''}".lower()
    if "baseline" in text and "300" in text:
        return ["mambayolo_t_ssdd_300e_16bs"]
    if "baseline" in text and "100" in text:
        return ["mambayolo_t_ssdd"]
    m = re.search(r"ex\s*([0-9]+)", text)
    if m:
        n = m.group(1)
        return [f"ex{n}", f"ex{n}-", f"ex{n}_"]
    if "pfim" in text:
        return ["pfim"]
    if "ds-ssm" in text or "dsssm" in text or "完整" in text:
        return ["ds-ssm"]
    return [text.strip()] if text.strip() else []


def find_match(row, ranking):
    raw = row.get("实验名_raw")
    exp_id = row.get("实验编号")
    queries = manual_to_query(raw, exp_id)
    candidates = []
    for r in ranking:
        name = r.get("experiment_name", "")
        nn = norm(name)
        ok = False
        for q in queries:
            nq = norm(q)
            if nq and nq in nn:
                ok = True
                break
        if not ok:
            continue
        if "ssdd" not in name.lower():
            continue
        candidates.append(r)
    if not candidates:
        return None, "未匹配"
    candidates.sort(key=lambda r: (r.get("score") is not None, r.get("score") or -1, r.get("training_complete")), reverse=True)
    best = candidates[0]
    reason = "唯一匹配" if len(candidates) == 1 else f"多候选取最高分({len(candidates)}个)"
    return best, reason


def copy_sheet(wb, src_name, dst_name):
    src = wb[src_name]
    if dst_name in wb.sheetnames:
        del wb[dst_name]
    dst = wb.copy_worksheet(src)
    dst.title = dst_name
    return dst


def main():
    os.makedirs(os.path.dirname(OUTPUT), exist_ok=True)
    wb = openpyxl.load_workbook(INPUT)
    ranking = json.load(open(RANKING, encoding="utf-8"))

    ws = copy_sheet(wb, "结构化总表", "结构化总表_核验更新")
    headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    extra_headers = ["实际来源服务器", "实际实验路径", "实际rank", "实际score", "匹配状态", "核验备注"]
    for h in extra_headers:
        if h not in headers:
            col = ws.max_column + 1
            ws.cell(1, col).value = h
            headers[h] = col

    yellow = PatternFill("solid", fgColor="FFF2CC")
    green = PatternFill("solid", fgColor="D9EAD3")
    red = PatternFill("solid", fgColor="F4CCCC")

    match_rows = []
    for r in range(2, ws.max_row + 1):
        row = {h: ws.cell(r, c).value for h, c in headers.items() if h}
        match, reason = find_match(row, ranking)
        if match:
            metrics = score_fields(match.get("metrics", {}))
            changed = []
            for field in ["P", "R", "mAP50", "mAP50-95"]:
                val = metrics.get(field)
                col = headers.get(field)
                if val is not None and col:
                    old = ws.cell(r, col).value
                    ws.cell(r, col).value = round(float(val), 6)
                    if old is not None and abs(float(old) - float(val)) > 1e-6:
                        ws.cell(r, col).fill = yellow
                        changed.append(f"{field}: {old} -> {round(float(val), 6)}")
            ws.cell(r, headers["实际来源服务器"]).value = match.get("source_server")
            ws.cell(r, headers["实际实验路径"]).value = match.get("remote_dir")
            ws.cell(r, headers["实际rank"]).value = match.get("rank")
            ws.cell(r, headers["实际score"]).value = match.get("score")
            ws.cell(r, headers["匹配状态"]).value = reason
            ws.cell(r, headers["核验备注"]).value = "; ".join(changed) if changed else "指标与可解析结果一致或原表缺失"
            ws.cell(r, headers["匹配状态"]).fill = green
            match_rows.append((r, row.get("实验名_raw"), match, reason, changed))
        else:
            ws.cell(r, headers["匹配状态"]).value = reason
            ws.cell(r, headers["核验备注"]).value = "未在本地实验索引/排名中找到明确对应项，保留原表记录"
            ws.cell(r, headers["匹配状态"]).fill = red

    if "核验说明" in wb.sheetnames:
        del wb["核验说明"]
    note = wb.create_sheet("核验说明", 0)
    note.append(["项目", "值"])
    note.append(["生成时间", datetime.now().isoformat(timespec="seconds")])
    note.append(["原始文件", INPUT])
    note.append(["新文件", OUTPUT])
    note.append(["依据", "database/experiment_ranking.json; recovered results.csv metadata"])
    note.append(["处理原则", "能明确匹配则更新实际指标/路径/rank；不能明确匹配则保留并标注未匹配"])
    note.append(["匹配成功行数", len(match_rows)])
    note.append(["总记录行数", ws.max_row - 1])

    if "匹配明细" in wb.sheetnames:
        del wb["匹配明细"]
    detail = wb.create_sheet("匹配明细")
    detail.append(["源表行", "手工实验名", "匹配服务器", "匹配实验路径", "rank", "score", "匹配说明", "修改摘要"])
    for source_row, manual_name, match, reason, changed in match_rows:
        detail.append([
            source_row,
            manual_name,
            match.get("source_server"),
            match.get("remote_dir"),
            match.get("rank"),
            match.get("score"),
            reason,
            "; ".join(changed),
        ])

    for sheet in [ws, note, detail]:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for col in range(1, min(sheet.max_column, 20) + 1):
            sheet.column_dimensions[get_column_letter(col)].width = min(max(len(str(sheet.cell(1, col).value or "")) + 4, 12), 36)
        sheet.freeze_panes = "A2"

    wb.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
