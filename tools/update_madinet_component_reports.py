from pathlib import Path
import csv
import json
import sys

sys.path.insert(0, r"C:\Users\13207\.cache\codex-runtimes\codex-primary-runtime\dependencies\python")

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path.cwd()
SUITE = ROOT / "downloads" / "complexity_results_172_26_151_143" / "madinet_component_suite_20260614_172647" / "madinet_component_suite.csv"


def read_csv(path):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def set_font(run, name="宋体"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def heading(doc, text, level=1):
    p = doc.add_heading("", level=level)
    r = p.add_run(text)
    set_font(r, "黑体")


def shade(row):
    for cell in row.cells:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "D9EAF7")
        cell._tc.get_or_add_tcPr().append(shd)


def main():
    rows = read_csv(SUITE)
    out_csv = ROOT / "database" / "madinet_official_component_suite.csv"
    out_csv.parent.mkdir(parents=True, exist_ok=True)
    with out_csv.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=rows[0].keys())
        writer.writeheader()
        writer.writerows(rows)

    md = ROOT / "reports" / "madinet_official_component_suite_report.md"
    lines = [
        "# MaDiNet 官方代码其它模块复杂度统计",
        "",
        "生成时间：2026-06-14",
        "",
        "服务器：`172.26.151.143`",
        "",
        "说明：这些结果来自 MaDiNet 官方仓库中的 `MambaSAR.py` 和 `pixel difference convPDC.py`，属于组件级统计，不是完整检测器 end-to-end 统计。",
        "",
        "| 组件 | 输入 | Params(M) | GFLOPs(THOP) | latency(ms) | 说明 |",
        "|---|---|---:|---:|---:|---|",
    ]
    for r in rows:
        gf = r["GFLOPs_thop"]
        if gf == "0.0" and r["component"].startswith("PDC"):
            gf = "THOP unsupported"
        lines.append(f"| {r['component']} | {r['input_shape']} | {r['params_M']} | {gf} | {r['latency_ms']} | {r['note']} |")
    lines += [
        "",
        "## 观察",
        "",
        "1. MaDiNet 官方 MambaVisionMixer 单块在 `d=1024, L=1600` 下约 2.18M 参数、6.95 GFLOPs。",
        "2. 官方 Block-Mamba 单块约 10.58M 参数、33.82 GFLOPs；Block-Attention 单块约 12.60M 参数、40.29 GFLOPs。",
        "3. 官方源码注释中的 res4+res5 组合包含多层 Block，因此整体核心模块达到 223.26M 参数、345.28 GFLOPs。",
        "4. PDC 方法参数量不大，但 THOP 无法识别函数式差分卷积 FLOPs；当前保留服务器实测 latency。",
        "",
        "## 对本文的启发",
        "",
        "MaDiNet 的 Mamba/Attention 模块具有较强表达能力，但单模块和多层堆叠的复杂度明显较高。本文 DS-SSM 的优势是将 SAR 去斑选择性建模压缩到完整约 6M 参数、约 13G FLOPs 的检测器中，适合写作成轻量高效路线。",
    ]
    md.write_text("\n".join(lines), encoding="utf-8")

    doc_path = ROOT / "paper_draft" / "DS_SSM_方法优势与MaDiNet复杂度实测对比_中文.docx"
    doc = Document(str(doc_path))
    heading(doc, "六、MaDiNet 官方其它模块统计补充", 1)
    doc.add_paragraph("除官方 MambaSAR res4+res5 核心模块外，我进一步统计了 MaDiNet 官方代码中出现的 MambaVisionMixer、Self-Attention、AgentAttention、Block 和 PDC 等组件。结果如下。")
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["组件", "Params(M)", "GFLOPs", "延迟(ms)", "说明"]):
        table.rows[0].cells[i].text = h
    shade(table.rows[0])
    for r in rows:
        cells = table.add_row().cells
        gf = r["GFLOPs_thop"]
        if gf == "0.0" and r["component"].startswith("PDC"):
            gf = "THOP未识别"
        vals = [r["component"], r["params_M"], gf, r["latency_ms"], r["note"]]
        for i, v in enumerate(vals):
            cells[i].text = str(v)
    doc.add_paragraph("这些组件结果进一步说明：MaDiNet 的核心模块表达能力强，但复杂度较高；本文 DS-SSM 更适合作为轻量高效 SAR 舰船检测方案来写。")
    doc.save(doc_path)

    print(json.dumps({"csv": str(out_csv), "report": str(md), "docx": str(doc_path)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
